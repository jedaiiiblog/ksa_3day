from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from .calc import Contribution
from .collector import MemberMonth
from .summary import Summaries

_FONT = "맑은 고딕"
_TITLE_ONLY = 5


def _style(text_frame, size, bold=False):
    for p in text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = _FONT


def _fill(cell, text, size, bold=False):
    cell.text_frame.word_wrap = True
    cell.text_frame.text = text
    _style(cell.text_frame, size, bold)


def _table(slide, rows, widths, top, size):
    shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(0.5), Inches(top), Inches(sum(widths)), Inches(0.4 * len(rows))
    )
    for i, w in enumerate(widths):
        shape.table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            _fill(shape.table.cell(r, c), str(value), size, bold=(r == 0))
    return shape


def _note(slide, text, top, size=12):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(1.0))
    box.text_frame.word_wrap = True
    box.text_frame.text = text
    _style(box.text_frame, size)


def _new_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[_TITLE_ONLY])
    slide.shapes.title.text = title
    _style(slide.shapes.title.text_frame, 28, bold=True)
    return slide


def _pct(v):
    return f"{v:.1f}%"


def _hours(v):
    return f"{v:.1f}h"


def build_deck(out_path, year: int, month: int, members: list[MemberMonth],
               contribs: dict[str, Contribution], summaries: Summaries) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    team = _new_slide(prs, f"{year}년 {month}월 팀 주간 보고 취합")
    rows = [["팀원", "프로젝트 수", "월 기준 시간", "팀 대비 비중"]]
    for mm in members:
        c = contribs[mm.member]
        rows.append([mm.member, c.project_count, _hours(c.base_hours), _pct(c.team_share_pct)])
    _table(team, rows, [3.5, 2.6, 3.1, 3.1], 1.5, 14)
    _note(team, summaries.team or "(AI 요약 없음)", 1.5 + 0.4 * len(rows) + 0.4, 14)

    for mm in members:
        c = contribs[mm.member]
        slide = _new_slide(prs, f"{mm.member} — {year}년 {month}월")
        rows = [["프로젝트", "요약", "시간", "개인 기준 %", "팀 대비 %"]]
        for name, bodies in mm.projects.items():
            rows.append([
                name,
                summaries.get(mm.member, name, bodies),
                _hours(c.per_project_hours),
                _pct(c.personal_pct),
                _pct(c.project_share_pct),
            ])
        size = 12 if len(rows) <= 6 else 9
        _table(slide, rows, [2.4, 5.6, 1.1, 1.6, 1.6], 1.4, size)
        memo = f"근무일 {c.workdays}일 · 월 기준 시간 {_hours(c.base_hours)} · 팀 대비 {_pct(c.team_share_pct)}"
        late = [Path(r.path).name for r in mm.reports if r.late]
        if late:
            memo += "\n마감 후 저장된 파일: " + ", ".join(late)
        _note(slide, memo, 6.5, 11)

    prs.save(str(out_path))
