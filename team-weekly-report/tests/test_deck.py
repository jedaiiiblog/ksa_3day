import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from pptx import Presentation

from team_report.calc import compute
from team_report.calendar_data import Calendar
from team_report.collector import MemberMonth
from team_report.deck import build_deck
from team_report.models import Project, Report
from team_report.summary import Summaries

HOL = {date(2026, 9, 24), date(2026, 9, 25)}


def slide_text(slide):
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if sh.has_table:
            for row in sh.table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


class DeckTest(unittest.TestCase):
    def build(self, members, summaries=None):
        counts = {m.member: len(m.projects) for m in members}
        contribs = compute(counts, 2026, 9, Calendar(holidays=set(HOL)))
        d = tempfile.TemporaryDirectory()
        self.addCleanup(d.cleanup)
        out = Path(d.name) / "out.pptx"
        build_deck(out, 2026, 9, members, contribs, summaries or Summaries(team="팀 요약"))
        return Presentation(str(out))

    def test_slide_count_order_and_numbers(self):
        a = MemberMonth("A", {f"P{i}": ["x"] for i in range(4)})
        b = MemberMonth("B", {"Q": ["y"]})
        prs = self.build([a, b])
        self.assertEqual(len(prs.slides), 3)
        team = slide_text(prs.slides[0])
        self.assertIn("2026년 9월", team)
        self.assertIn("팀 요약", team)
        self.assertIn("50.0%", team)  # A, B 모두 160h -> 각 50%
        sa = slide_text(prs.slides[1])
        self.assertIn("A — 2026년 9월", sa)
        self.assertIn("40.0h", sa)     # 160 / 4
        self.assertIn("25.0%", sa)     # 개인 기준
        self.assertIn("12.5%", sa)     # 40 / 320
        self.assertIn("근무일 20일", sa)
        self.assertIn("160.0h", sa)
        self.assertIn("B — 2026년 9월", slide_text(prs.slides[2]))

    def test_summary_used_and_fallback_marked(self):
        a = MemberMonth("A", {"P1": ["원문1"], "P2": ["원문2"]})
        s = Summaries(team="t", projects={("a", "p1"): "요약된 문장"})
        text = slide_text(self.build([a], s).slides[1])
        self.assertIn("요약된 문장", text)
        self.assertIn("(AI 요약 없음) 원문2", text)

    def test_late_files_listed_on_member_slide(self):
        r = Report("A", date(2026, 9, 25), [Project("P", "x")], path="C:/x/주간보고_a_1.md",
                   saved_at=datetime(2026, 9, 28, 9), late=True)
        a = MemberMonth("A", {"P": ["x"]}, [r])
        self.assertIn("주간보고_a_1.md", slide_text(self.build([a]).slides[1]))

    def test_member_with_many_projects_and_long_text_still_builds(self):
        a = MemberMonth("A", {f"프로젝트{i}": ["가" * 400] for i in range(12)})
        prs = self.build([a])
        self.assertEqual(len(prs.slides), 2)

    def test_no_members_builds_team_slide_only(self):
        self.assertEqual(len(self.build([]).slides), 1)


if __name__ == "__main__":
    unittest.main()
